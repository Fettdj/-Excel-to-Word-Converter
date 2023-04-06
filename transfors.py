import openpyxl
from docx import Document

def convertir_excel_a_word(archivo_excel, archivo_word):
    # Cargar el archivo de Excel
    workbook = openpyxl.load_workbook(archivo_excel)
    sheet = workbook.active

    # Crear un nuevo documento de Word
    doc = Document()

    # Iterar sobre las filas y columnas de la hoja de cálculo de Excel
    for row in sheet.iter_rows():
        fila = []
        for cell in row:
            fila.append(str(cell.value))

        # Agregar una nueva línea en el documento de Word con los datos de la fila
        doc.add_paragraph(', '.join(fila))

    # Guardar el documento de Word
    doc.save(archivo_word)

if __name__ == "__main__":
    archivo_excel = 'ejemplo.xlsx'  # Reemplaza esto con el nombre de tu archivo de Excel
    archivo_word = 'salida.docx'  # Reemplaza esto con el nombre que deseas para el archivo de Word

    convertir_excel_a_word(archivo_excel, archivo_word)
    print(f"Archivo de Word generado: {archivo_word}")
